"""File upload API for chat — saves files to agent workspace and extracts text."""

import base64
import os
import uuid
from pathlib import Path

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, Form
from app.core.security import get_current_user
from app.models.user import User
from app.config import get_settings

router = APIRouter(prefix="/chat", tags=["chat"])

_settings = get_settings()
WORKSPACE_ROOT = Path(_settings.AGENT_DATA_DIR)

# Supported extensions and their text extraction method
TEXT_EXTENSIONS = {
    ".txt", ".md", ".csv", ".json", ".xml", ".yaml", ".yml",
    ".py", ".js", ".ts", ".html", ".css", ".sql", ".sh", ".log",
    ".ini", ".cfg", ".conf", ".env", ".toml",
}
OFFICE_EXTENSIONS = {".pdf", ".docx", ".doc", ".xlsx", ".xls", ".pptx", ".ppt"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}
EXTRACTABLE = TEXT_EXTENSIONS | OFFICE_EXTENSIONS

MIME_MAP = {
    ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
    ".gif": "image/gif", ".webp": "image/webp", ".bmp": "image/bmp",
}


def extract_text(file_path: Path, extension: str) -> str:
    """Extract text content from a file."""
    if extension in TEXT_EXTENSIONS:
        try:
            return file_path.read_text(encoding="utf-8", errors="replace")
        except Exception:
            return file_path.read_text(encoding="gbk", errors="replace")

    if extension == ".pdf":
        try:
            import subprocess
            result = subprocess.run(
                ["python3", "-c", f"""
import sys
try:
    import PyPDF2
    reader = PyPDF2.PdfReader('{file_path}')
    text = '\\n'.join(page.extract_text() or '' for page in reader.pages)
    print(text[:8000])
except ImportError:
    # Fallback: use pdftotext if available
    import subprocess as sp
    r = sp.run(['pdftotext', '{file_path}', '-'], capture_output=True, text=True)
    print(r.stdout[:8000] if r.returncode == 0 else '[无法解析PDF]')
"""],
                capture_output=True, text=True, timeout=30,
            )
            return result.stdout.strip() or "[PDF内容提取失败]"
        except Exception as e:
            return f"[PDF解析错误: {e}]"

    if extension == ".docx":
        try:
            import subprocess
            result = subprocess.run(
                ["python3", "-c", f"""
try:
    from docx import Document
    doc = Document('{file_path}')
    text = '\\n'.join(p.text for p in doc.paragraphs)
    print(text[:8000])
except ImportError:
    print('[需要安装 python-docx 库]')
"""],
                capture_output=True, text=True, timeout=30,
            )
            return result.stdout.strip() or "[DOCX内容提取失败]"
        except Exception as e:
            return f"[DOCX解析错误: {e}]"

    if extension in (".xlsx", ".xls"):
        try:
            import subprocess
            result = subprocess.run(
                ["python3", "-c", f"""
try:
    import openpyxl
    wb = openpyxl.load_workbook('{file_path}', read_only=True)
    lines = []
    for ws in wb.worksheets[:3]:
        lines.append(f'## Sheet: {{ws.title}}')
        for row in ws.iter_rows(max_row=50, values_only=True):
            lines.append('\\t'.join(str(c) if c is not None else '' for c in row))
    print('\\n'.join(lines)[:8000])
except ImportError:
    print('[需要安装 openpyxl 库]')
"""],
                capture_output=True, text=True, timeout=30,
            )
            return result.stdout.strip() or "[Excel内容提取失败]"
        except Exception as e:
            return f"[Excel解析错误: {e}]"

    return f"[不支持的文件格式: {extension}]"


@router.post("/upload")
async def upload_file(
    file: UploadFile = File(...),
    agent_id: str = Form(""),
    current_user: User = Depends(get_current_user),
):
    """Upload a file for chat context. Saves to agent workspace/uploads/ and returns extracted text."""
    if not file.filename:
        raise HTTPException(status_code=400, detail="No filename")

    ext = os.path.splitext(file.filename)[1].lower()

    content = await file.read()

    # Determine save directory
    workspace_path = ""
    if agent_id:
        # Save to agent's workspace/uploads/
        uploads_dir = WORKSPACE_ROOT / agent_id / "workspace" / "uploads"
        uploads_dir.mkdir(parents=True, exist_ok=True)
        save_path = uploads_dir / file.filename
        # Avoid overwriting: add suffix if file exists
        if save_path.exists():
            stem = save_path.stem
            suffix = save_path.suffix
            counter = 1
            while save_path.exists():
                save_path = uploads_dir / f"{stem}_{counter}{suffix}"
                counter += 1
        save_path.write_bytes(content)
        workspace_path = f"workspace/uploads/{save_path.name}"
    else:
        # Fallback: save to /tmp (legacy behavior)
        fallback_dir = Path("/tmp/clawith_uploads")
        fallback_dir.mkdir(exist_ok=True)
        file_id = str(uuid.uuid4())[:8]
        save_path = fallback_dir / f"{file_id}_{file.filename}"
        save_path.write_bytes(content)

    # Extract text (only for known formats)
    is_image = ext in IMAGE_EXTENSIONS
    image_data_url = ""
    if is_image:
        # For images: generate base64 data URL for vision models
        if len(content) > 10 * 1024 * 1024:  # 10MB limit
            raise HTTPException(status_code=400, detail="Image too large (max 10MB)")
        mime = MIME_MAP.get(ext, "image/png")
        b64 = base64.b64encode(content).decode("ascii")
        image_data_url = f"data:{mime};base64,{b64}"
        extracted = f"[图片文件: {file.filename}，需要视觉模型分析]"
    elif ext in EXTRACTABLE:
        extracted = extract_text(save_path, ext)
    else:
        extracted = f"[文件已保存，格式 {ext} 暂不支持文本提取，Agent 可通过 read_document 工具读取]"

    # Truncate if too long
    if len(extracted) > 6000:
        extracted = extracted[:6000] + "\n\n...[内容已截断，共 " + str(len(extracted)) + " 字]"

    return {
        "filename": file.filename,
        "saved_filename": save_path.name,
        "size": len(content),
        "extracted_text": extracted,
        "workspace_path": workspace_path,
        "is_image": is_image,
        "image_data_url": image_data_url,
    }
